"""
Ядро сортировки. Понимает два формата выгрузок (определяет сам по колонкам):

  1) «по коробам»   — колонки Модель / Код упаковки SSCC / Размер (каждая пара отдельной строкой)
  2) «график поставок» — колонки Модель / Кол-во коробов / Кол-во пар (одна строка на модель)

Файлы можно загружать вперемешку: каждый формат собирается в свой общий файл
рядом с исходным: <имя>_sortirovka.xlsx

Здесь нет графики и нет pandas — только openpyxl. Так exe получается маленьким.
"""
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

ITEMS_PER_BOX = 8                 # сколько пар должно быть в полном коробе
SUFFIX = "_sortirovka"            # что добавляем к имени файла
NO_MODEL = "(без модели)"         # лист для строк, где модель не заполнена
KIND_NAMES = {"boxes": "по коробам (SSCC)", "schedule": "график поставок"}

# Названия колонок формата «график поставок» (в нижнем регистре)
SCHEDULE_FIELDS = {
    "model": "модель", "boxes": "кол-во коробов", "pairs": "кол-во пар",
    "date": "дата", "time": "время", "total_pairs": "кол-во пар общее",
    "total_boxes": "кол-во коробов общее", "volume": "объем поставки, куб",
    "kari": "артикул кари", "box_art": "артикул короба", "barcode": "шк короба",
    "price": "цена закупки", "note": "примеч",
}


class SortError(Exception):
    """Понятная пользователю ошибка (её текст показываем в окне)."""


# ---------- чтение ----------

def _norm(value):
    """Приводит заголовок к простому виду: без пробелов по краям, строчными буквами."""
    return str(value).strip().lower() if value is not None else ""


def _find_columns(header):
    """Ищет номера колонок: модель, SSCC, размер. Возвращает None, если чего-то нет."""
    names = [_norm(v) for v in header]
    model = next((i for i, n in enumerate(names) if n == "модель"), None)
    sscc = next((i for i, n in enumerate(names) if "sscc" in n), None)
    size = next((i for i, n in enumerate(names) if n == "размер"), None)
    if None in (model, sscc, size):
        return None
    return model, sscc, size


def _find_schedule_cols(header):
    """Колонки формата «график поставок»: словарь поле -> номер колонки (или None)."""
    names = [_norm(v) for v in header]
    cols = {key: (names.index(title) if title in names else None)
            for key, title in SCHEDULE_FIELDS.items()}
    if None in (cols["model"], cols["boxes"], cols["pairs"]):
        return None
    return cols


def read_file(path):
    """Читает файл и сам определяет формат. Возвращает словарь:
    kind ("boxes"/"schedule"), header, rows, cols."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise SortError("Не удалось открыть файл. Убедитесь, что это обычный Excel (.xlsx).")

    try:
        for ws in wb.worksheets:
            kind, header, cols, rows = None, None, None, []
            for row in ws.iter_rows(values_only=True):
                if header is None:
                    # ищем строку с заголовками (обычно это первая строка)
                    found = _find_columns(row)
                    if found:
                        kind, header, cols = "boxes", list(row), found
                    else:
                        found = _find_schedule_cols(row)
                        if found:
                            kind, header, cols = "schedule", list(row), found
                    continue
                if all(v is None or str(v).strip() == "" for v in row):
                    continue                      # пропускаем пустые строки
                rows.append(list(row) + [None] * (len(header) - len(row)))
            if header is not None:
                if not rows:
                    raise SortError("В файле есть заголовки, но нет ни одной строки с данными.")
                return {"kind": kind, "header": header, "rows": rows, "cols": cols}
    finally:
        wb.close()

    raise SortError("Не распознал формат. Нужны колонки «Модель», «Код упаковки SSCC», «Размер» "
                    "или «Модель», «Кол-во коробов», «Кол-во пар».")


# ---------- обработка ----------

def _size_key(value):
    """Ключ сортировки размеров: числа по возрастанию (9 раньше 10), потом текст."""
    try:
        return (0, float(str(value).replace(",", ".")), "")
    except ValueError:
        return (1, 0.0, str(value))


def group_rows(rows, cols):
    """Раскладывает строки по моделям и сортирует: короб -> размер."""
    c_model, c_sscc, c_size = cols
    groups = defaultdict(list)
    for row in rows:
        model = str(row[c_model]).strip() if row[c_model] is not None else ""
        groups[model or NO_MODEL].append(row)

    for model_rows in groups.values():
        model_rows.sort(key=lambda r: (str(r[c_sscc]), _size_key(r[c_size])))
    return dict(sorted(groups.items()))


def check_boxes(rows, cols):
    """Считает проблемные короба: больше 8 пар или разные модели в одном коробе."""
    c_model, c_sscc, _ = cols
    items = defaultdict(int)          # SSCC -> сколько пар
    models = defaultdict(set)         # SSCC -> какие модели
    for row in rows:
        if row[c_sscc] is None or str(row[c_sscc]).strip() == "":
            continue
        sscc = str(row[c_sscc]).strip()
        items[sscc] += 1
        models[sscc].add(str(row[c_model]).strip() if row[c_model] is not None else "")

    bad = [s for s in items if items[s] > ITEMS_PER_BOX or len(models[s]) > 1]
    incomplete = [s for s in items if items[s] < ITEMS_PER_BOX]
    return len(items), len(bad), len(incomplete)


# ---------- запись ----------

def _clean(value):
    """Спецсимволы (например GS в кодах маркировки) записываем в виде _x001D_,
    как это делает сам Excel, — так данные не теряются и не ломают запись."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub(lambda m: "_x%04X_" % ord(m.group()), value)
    return value


def _safe_sheet_name(name, used):
    """Имя листа: до 31 символа, без []:*?/\\ и без повторов."""
    name = re.sub(r"[\[\]:*?/\\]", "_", name).strip("'") or "Лист"
    name = name[:31]
    base, i = name, 2
    while name.lower() in used:
        suffix = f"_{i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name


def _style_sheet(ws):
    """Оформление: цветная шапка, ширина колонок, закреплённая первая строка."""
    fill = PatternFill("solid", fgColor="2B3140")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for col in ws.columns:
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(10, min(longest + 3, 60))
    ws.freeze_panes = "A2"


def _write_summary(ws, groups, cols, file_stats):
    """Лист «Сводка»: модель | коробов | пар | по размерам | итого.
    Если файлов несколько — ниже добавляем список, сколько пар было в каждом."""
    c_sscc, c_size = cols[1], cols[2]
    all_sizes = sorted({str(r[c_size]) for g in groups.values() for r in g}, key=_size_key)

    ws.append(["Модель", "Коробов", "Пар"] + all_sizes)
    totals = [0, 0] + [0] * len(all_sizes)

    for model, model_rows in groups.items():
        boxes = len({str(r[c_sscc]) for r in model_rows})
        per_size = [sum(1 for r in model_rows if str(r[c_size]) == s) for s in all_sizes]
        line = [boxes, len(model_rows)] + per_size
        ws.append([model] + line)
        totals = [a + b for a, b in zip(totals, line)]

    ws.append(["ИТОГО"] + totals)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    if len(file_stats) > 1:
        ws.append([])
        ws.append(["Обработанные файлы", "Пар"])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for name, pairs in file_stats:
            ws.append([name, pairs])

    _style_sheet(ws)


def output_path_for(paths):
    """Один файл:  MC_185107.xlsx -> MC_185107_sortirovka.xlsx
    Несколько:    MC_185107_MC_185108_sortirovka.xlsx (в папке первого файла)."""
    paths = [Path(p) for p in paths]
    stems = [p.stem for p in paths]
    name = "_".join(stems)
    if len(name) > 80:                       # слишком длинное имя — сокращаем
        name = f"{stems[0]}_plus{len(stems) - 1}"
    return paths[0].with_name(name + SUFFIX + ".xlsx")


def _save(wb, out):
    try:
        wb.save(out)
    except PermissionError:
        raise SortError(f"Не удалось сохранить «{out.name}». Если он открыт в Excel — закройте его и повторите.")
    except OSError:
        raise SortError("Не удалось сохранить файл рядом с исходным. Проверьте доступ к папке.")


# ---------- формат 1: «по коробам» ----------

def _process_boxes(items):
    """items — список (путь, таблица). Объединяет все файлы в один результат."""
    header, cols = items[0][1]["header"], items[0][1]["cols"]
    all_rows, file_stats = [], []
    for path, table in items:
        if [_norm(v) for v in table["header"]] != [_norm(v) for v in header]:
            # колонки отличаются — молча склеивать нельзя, данные съедут
            raise SortError(f"{path.name}: колонки не совпадают с первым файлом "
                            f"({items[0][0].name}). Все файлы одного формата должны быть одинаковыми.")
        all_rows.extend(table["rows"])
        file_stats.append((path.name, len(table["rows"])))

    groups = group_rows(all_rows, cols)
    boxes, bad, incomplete = check_boxes(all_rows, cols)

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Сводка"
    _write_summary(summary_ws, groups, cols, file_stats)

    used = {"сводка"}
    for model, model_rows in groups.items():
        ws = wb.create_sheet(_safe_sheet_name(model, used))
        ws.append([_clean(h) for h in header])
        for row in model_rows:
            ws.append([_clean(v) for v in row])
        _style_sheet(ws)

    out = output_path_for([p for p, _ in items])
    _save(wb, out)
    warnings = []
    if bad:
        warnings.append(f"В {bad} коробах больше {ITEMS_PER_BOX} пар или смешаны разные модели — проверьте исходник")
    if incomplete:
        warnings.append(f"Неполных коробов (меньше {ITEMS_PER_BOX} пар): {incomplete}")
    return {"kind": "boxes", "files": len(items), "output": str(out), "models": len(groups),
            "boxes": boxes, "pairs": len(all_rows), "warnings": warnings, "warn_level": "bad" if bad else "info"}


# ---------- формат 2: «график поставок» ----------

def _num(value):
    try:
        return int(float(str(value).replace(",", ".").replace(" ", "")))
    except ValueError:
        return 0


def _to_date(value):
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _read_deliveries(items):
    """Из таблиц собирает поставки (по одной на файл) и строки по моделям."""
    deliveries, lines, warnings, used_keys = [], [], [], set()
    for path, table in items:
        c, rows = table["cols"], table["rows"]
        get = lambda r, k: r[c[k]] if c[k] is not None else None
        first = next((r for r in rows if get(r, "date") is not None), rows[0])
        date = _to_date(get(first, "date"))

        stem = path.stem
        city = re.split(r"[_\s]", stem)[0]
        if city[:1].isdigit():
            city = ""                                    # имя начинается с цифры — города нет
        key = f"{city} {date:%d.%m.%Y}".strip() if date else stem
        while key in used_keys:
            key += "*"
        used_keys.add(key)

        # дата в названии файла и внутри файла должны совпадать
        m = re.search(r"_(\d\d)_(\d\d)_(\d{2,4})", stem)
        if m and date:
            d, mo, y = int(m[1]), int(m[2]), int(m[3])
            y = y + 2000 if y < 100 else y
            if (d, mo, y) != (date.day, date.month, date.year):
                warnings.append(f"{path.name}: в названии {d:02d}.{mo:02d}.{y}, "
                                f"а внутри файла {date:%d.%m.%Y} — взята дата из файла")

        sum_pairs = sum_boxes = odd = 0
        for r in rows:
            model = str(get(r, "model")).strip() if get(r, "model") is not None else ""
            if not model:
                continue
            boxes, pairs = _num(get(r, "boxes")), _num(get(r, "pairs"))
            sum_boxes += boxes
            sum_pairs += pairs
            odd += pairs != boxes * ITEMS_PER_BOX
            lines.append({"key": key, "date": date, "city": city, "box_art": get(r, "box_art"),
                          "barcode": get(r, "barcode"), "kari": get(r, "kari"), "model": model,
                          "price": get(r, "price"), "boxes": boxes, "pairs": pairs,
                          "note": get(r, "note"), "file": path.name})

        decl_pairs, decl_boxes = get(first, "total_pairs"), get(first, "total_boxes")
        if decl_pairs is not None and _num(decl_pairs) != sum_pairs:
            warnings.append(f"{path.name}: пар по моделям {sum_pairs}, в шапке файла {_num(decl_pairs)}")
        if decl_boxes is not None and _num(decl_boxes) != sum_boxes:
            warnings.append(f"{path.name}: коробов по моделям {sum_boxes}, в шапке файла {_num(decl_boxes)}")
        if odd:
            warnings.append(f"{path.name}: строк, где пар не равно коробам × {ITEMS_PER_BOX}: {odd}")

        deliveries.append({"key": key, "date": date, "time": get(first, "time"), "city": city,
                           "file": path.name, "pairs": decl_pairs, "boxes": decl_boxes,
                           "volume": get(first, "volume"), "sum_pairs": sum_pairs})
    deliveries.sort(key=lambda d: (d["date"] is None, d["date"] or datetime.min, d["key"]))
    return deliveries, lines, warnings


def _bold_row(ws, row):
    for cell in ws[row]:
        cell.font = Font(bold=True)


def _process_schedule(items):
    deliveries, lines, warnings = _read_deliveries(items)
    by_model = defaultdict(list)
    for ln in lines:
        by_model[ln["model"]].append(ln)
    models = sorted(by_model)
    for m in models:
        by_model[m].sort(key=lambda ln: (ln["date"] is None, ln["date"] or datetime.min))

    # артикул и цена по модели должны быть одни и те же во всех поставках
    for m in models:
        for field, title in (("kari", "артикулов"), ("price", "цен закупки")):
            values = {ln[field] for ln in by_model[m] if ln[field] is not None}
            if len(values) > 1:
                warnings.append(f"{m}: несколько {title}: {', '.join(str(v) for v in sorted(values, key=str))}")

    wb = Workbook()
    sm = wb.active
    sm.title = "Сводка"
    dl = wb.create_sheet("Поставки")

    # листы моделей
    model_cols = ["Поставка", "Дата", "Город", "Артикул короба", "ШК короба", "Артикул Кари",
                  "Модель", "Цена закупки", "Коробов", "Пар", "Примеч", "Файл"]
    used = {"сводка", "поставки"}
    for m in models:
        ws = wb.create_sheet(_safe_sheet_name(m, used))
        ws.append(model_cols)
        for ln in by_model[m]:
            ws.append([ln["key"], ln["date"], ln["city"], ln["box_art"],
                       None if ln["barcode"] is None else str(ln["barcode"]), ln["kari"], ln["model"],
                       ln["price"], ln["boxes"], ln["pairs"], ln["note"], ln["file"]])
            ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY"
        ws.append([])
        ws.append(["ИТОГО"] + [None] * 7 + [sum(ln["boxes"] for ln in by_model[m]),
                                           sum(ln["pairs"] for ln in by_model[m])])
        _style_sheet(ws)
        _bold_row(ws, ws.max_row)

    # сводка: модель | артикул | коробов | пар | поставок | пары по каждой поставке
    keys = [d["key"] for d in deliveries]
    sm.append(["Модель", "Артикул Кари", "Коробов", "Пар", "Поставок"] + keys)
    per_key = {k: 0 for k in keys}
    tot_boxes = tot_pairs = 0
    for m in models:
        mine = by_model[m]
        kari = next((ln["kari"] for ln in mine if ln["kari"] is not None), None)
        cells = [sum(ln["pairs"] for ln in mine if ln["key"] == k) for k in keys]
        for k, v in zip(keys, cells):
            per_key[k] += v
        boxes, pairs = sum(ln["boxes"] for ln in mine), sum(ln["pairs"] for ln in mine)
        tot_boxes += boxes
        tot_pairs += pairs
        sm.append([m, kari, boxes, pairs, len({ln["key"] for ln in mine})] + cells)
    sm.append(["ИТОГО", None, tot_boxes, tot_pairs, None] + [per_key[k] for k in keys])
    _style_sheet(sm)
    _bold_row(sm, sm.max_row)
    sm.append([])
    sm.append(["Колонки с названиями поставок (город и дата) — количество пар в этой поставке."])

    # лист «Поставки»: что заявлено в шапке файла и что получилось по моделям
    dl.append(["Поставка", "Дата", "Время", "Город", "Файл", "Пар по шапке файла",
               "Коробов по шапке файла", "Объём, куб", "Пар по моделям"])
    for d in deliveries:
        dl.append([d["key"], d["date"], d["time"], d["city"], d["file"],
                   None if d["pairs"] is None else _num(d["pairs"]),
                   None if d["boxes"] is None else _num(d["boxes"]),
                   d["volume"], d["sum_pairs"]])
        dl.cell(dl.max_row, 2).number_format = "DD.MM.YYYY"
    _style_sheet(dl)
    dl.row_dimensions[1].height = 32
    dl.append(["ИТОГО", None, None, None, None] +
              [sum(v for v in (d[f] and _num(d[f]) for d in deliveries) if v) for f in ("pairs", "boxes")] +
              [None, sum(d["sum_pairs"] for d in deliveries)])
    _bold_row(dl, dl.max_row)
    if warnings:
        dl.append([])
        dl.append(["Замечания"])
        _bold_row(dl, dl.max_row)
        for w in warnings:
            dl.append(["• " + w])

    out = output_path_for([p for p, _ in items])
    _save(wb, out)
    return {"kind": "schedule", "files": len(items), "output": str(out), "models": len(models),
            "boxes": tot_boxes, "pairs": tot_pairs, "warnings": warnings, "warn_level": "bad"}


# ---------- главная функция ----------

def process_files(paths, progress=None):
    """Читает файлы, сам определяет формат каждого и собирает по одному общему файлу
    на каждый формат. Возвращает список результатов (обычно один).
    progress(имя_файла, номер, всего) — необязательная функция для показа прогресса."""
    paths = [Path(p) for p in paths]
    if not paths:
        raise SortError("Не выбрано ни одного файла.")

    groups = {"boxes": [], "schedule": []}
    for i, path in enumerate(paths, 1):
        if progress:
            progress(path.name, i, len(paths))
        try:
            table = read_file(path)
        except SortError as e:
            raise SortError(f"{path.name}: {e}")
        groups[table["kind"]].append((path, table))

    results = []
    if groups["boxes"]:
        results.append(_process_boxes(groups["boxes"]))
    if groups["schedule"]:
        results.append(_process_schedule(groups["schedule"]))
    return results
